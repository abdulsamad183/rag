from __future__ import annotations

import io

from docx import Document as DocxDocument

from app.core.errors import IngestionError
from app.parsers.base import (
    BlockType,
    DocumentParser,
    ParsedBlock,
    ParsedDocument,
    render_table_markdown,
)
from app.utils.text import clean_text


class DocxParser(DocumentParser):
    source_types = ("docx",)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        try:
            doc = DocxDocument(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise IngestionError(f"Could not read DOCX: {exc}") from exc

        blocks: list[ParsedBlock] = []
        stack: list[tuple[int, str]] = []
        title = ""

        def section_path() -> str:
            return " > ".join(h for _, h in stack)

        for paragraph in doc.paragraphs:
            text = clean_text(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style.startswith("heading"):
                try:
                    level = int(style.replace("heading", "").strip() or "1")
                except ValueError:
                    level = 1
                if not title and level == 1:
                    title = text
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
                blocks.append(
                    ParsedBlock(
                        type=BlockType.HEADING, text=text, heading=text,
                        section_path=section_path(), heading_level=level,
                    )
                )
            elif style == "title" and not title:
                title = text
                blocks.append(ParsedBlock(type=BlockType.HEADING, text=text, heading=text, heading_level=0))
            else:
                blocks.append(
                    ParsedBlock(
                        type=BlockType.TEXT, text=text,
                        heading=stack[-1][1] if stack else "",
                        section_path=section_path(),
                    )
                )

        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers, body = rows[0], rows[1:]
            blocks.append(
                ParsedBlock(
                    type=BlockType.TABLE,
                    text=render_table_markdown(headers, body),
                    heading=stack[-1][1] if stack else "",
                    section_path=section_path(),
                    meta={"table": {"headers": headers, "rows": body}},
                )
            )

        core = doc.core_properties
        return ParsedDocument(
            title=title or (core.title or "") or filename,
            source_type="docx",
            blocks=blocks,
            meta={"author": core.author or ""},
        )
