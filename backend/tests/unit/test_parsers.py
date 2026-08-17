import json

import pytest

from app.core.errors import IngestionError
from app.parsers import get_parser
from app.security.files import detect_source_type


def parse(filename: str, payload: bytes):
    source_type = detect_source_type(filename)
    parser = get_parser(source_type)
    return parser.parse(payload, filename)


def full_text(doc) -> str:
    return " ".join(b.text for b in doc.blocks)


def test_txt_parser_paragraphs():
    doc = parse("notes.txt", b"first paragraph\n\nsecond paragraph")
    assert doc.source_type == "txt"
    assert len(doc.blocks) == 2


def test_markdown_parser_headings_and_code():
    content = b"""# Title

Intro text.

## Section A

Body of section A.

```python
print("hello")
```
"""
    doc = parse("guide.md", content)
    paths = [b.section_path for b in doc.blocks]
    assert any("Section A" in p for p in paths)
    code_blocks = [b for b in doc.blocks if b.type == "code"]
    assert len(code_blocks) == 1
    assert 'print("hello")' in code_blocks[0].text


def test_html_parser_strips_scripts():
    html = b"""<html><head><title>Page Title</title><script>evil()</script></head>
    <body><h1>Heading</h1><p>Visible text.</p><nav>menu</nav></body></html>"""
    doc = parse("page.html", html)
    text = full_text(doc)
    assert "Visible text." in text
    assert "evil" not in text
    assert "menu" not in text
    assert doc.title == "Page Title"


def test_csv_parser_keeps_headers():
    csv_data = b"name,latency\naurora,9ms\nbaseline,18ms\n"
    doc = parse("data.csv", csv_data)
    assert doc.blocks
    assert "name" in doc.blocks[0].text
    assert "aurora" in doc.blocks[0].text


def test_json_parser_flattens():
    payload = json.dumps({"system": {"name": "Aurora", "recall": 0.93}}).encode()
    doc = parse("config.json", payload)
    text = full_text(doc)
    assert "Aurora" in text
    assert "0.93" in text


def test_jsonl_parser_one_block_per_line():
    payload = b'{"q": "one"}\n{"q": "two"}\n'
    doc = parse("data.jsonl", payload)
    assert len(doc.blocks) == 2


def make_text_pdf(text: str) -> bytes:
    """Assemble a minimal but structurally valid single-page PDF with text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def test_pdf_parser_extracts_text():
    doc = parse("report.pdf", make_text_pdf("Aurora achieved 93 percent recall."))
    assert doc.source_type == "pdf"
    assert "93 percent recall" in full_text(doc)
    assert doc.blocks[0].page == 1


def test_pdf_parser_rejects_textless(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "blank.pdf"
    with open(path, "wb") as fh:
        writer.write(fh)

    with pytest.raises(IngestionError):
        parse("blank.pdf", path.read_bytes())


def test_docx_parser(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("Report Title", level=1)
    document.add_paragraph("Aurora achieved 93 percent recall.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "metric"
    table.rows[0].cells[1].text = "value"
    table.rows[1].cells[0].text = "recall"
    table.rows[1].cells[1].text = "0.93"
    path = tmp_path / "report.docx"
    document.save(path)

    doc = parse("report.docx", path.read_bytes())
    text = full_text(doc)
    assert "93 percent recall" in text
    table_blocks = [b for b in doc.blocks if b.type == "table"]
    assert table_blocks and "recall" in table_blocks[0].text
